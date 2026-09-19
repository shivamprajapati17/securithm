"""Build a NATIVE, fully-editable PPTX of the AuditAI deck.

Mirrors scripts/deck_draft.py layouts (1440x810 px -> 13.333x7.5 in slides),
using the same fonts as the reference deck: Josefin Sans (headings) + Poppins (body).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

S = 13.333 / 1440.0  # px -> inches

BG = RGBColor(0x05, 0x05, 0x08)
PANEL = RGBColor(0x0B, 0x0B, 0x10)
PANEL2 = RGBColor(0x11, 0x11, 0x1A)
BORDER = RGBColor(0x39, 0x47, 0xFF)
PURPLE = RGBColor(0xD8, 0x93, 0xF2)
YELLOW = RGBColor(0xFF, 0xFF, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xBD, 0xB7, 0xD0)
GREEN = RGBColor(0x2E, 0xE6, 0xA8)
RED = RGBColor(0xFF, 0x4D, 0x5A)
ORANGE = RGBColor(0xFF, 0x96, 0x3C)

HEAD = 'Josefin Sans'
BODY = 'Poppins'


def IN(v):
    return Inches(v * S)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def rect(slide, x, y, w, h, fill=PANEL, line=None, lw=1.25, rad=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    try:
        shp.adjustments[0] = rad
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def vline(slide, x, y, h, color=BORDER, weight=3.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(x), IN(y), Emu(9525 * weight), IN(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def dot(slide, x, y, w, h, fill=BORDER, line=WHITE, lw=1.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, IN(x), IN(y), IN(w), IN(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def arrow(slide, x, y, w, h, fill=BORDER):
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, IN(x), IN(y), IN(w), IN(h))
    shp.rotation = 90
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def hline(slide, x, y, w, color=BORDER, weight=3.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(x), IN(y), IN(w), Emu(9525 * weight))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def txt(slide, x, y, w, h, s, sz, color=WHITE, bold=False, font=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = s.split('\n')
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def frame(slide, title, kicker):
    rect(slide, 78, 83, 1282, 672, fill=None, line=BORDER, lw=2.2, rad=0.04)
    rect(slide, 1010, 22, 348, 84, fill=PANEL, line=None)
    txt(slide, 1010, 22, 348, 84, 'AUDITAI', 26, WHITE, True, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(slide, 110, 96, 600, 22, kicker, 11, YELLOW, True)
    txt(slide, 110, 120, 1100, 46, title, 24, WHITE, True, HEAD)


def wrap(s, limit):
    words, lines, cur = s.split(), [], ''
    for w in words:
        if len(cur + ' ' + w) <= limit:
            cur = (cur + ' ' + w).strip()
        else:
            lines.append(cur)
            cur = w
    lines.append(cur)
    return lines


def page1(prs):
    s = new_slide(prs)
    rect(s, 60, 50, 1320, 710, fill=None, line=BORDER, lw=2.2, rad=0.04)
    rect(s, 520, 92, 400, 78, fill=PANEL, line=BORDER, lw=1.5)
    txt(s, 520, 92, 400, 78, 'AUDITAI', 28, PURPLE, True, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, 320, 212, 800, 26, 'PRESENTS', 12, GREY, False, BODY, PP_ALIGN.CENTER)
    txt(s, 220, 262, 1000, 90, 'AUDIT AI', 60, WHITE, True, HEAD, PP_ALIGN.CENTER)
    txt(s, 220, 388, 1000, 36, 'AI-Powered Smart Contract Security Platform', 20, YELLOW, True, HEAD, PP_ALIGN.CENTER)
    txt(s, 220, 443, 1000, 24, 'Static Analysis  •  Symbolic Execution  •  AI Detection  •  Continuous Monitoring', 11, GREY, False, BODY, PP_ALIGN.CENTER)
    rect(s, 460, 520, 520, 62, fill=PANEL, line=BORDER, lw=1.5)
    txt(s, 460, 520, 520, 62, 'Team Unfiltered Devs   |   Securithm Product Line', 14, WHITE, True, BODY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, 320, 648, 800, 24, '2026  •  SaaS Security Audit Platform  •  auditai.io', 11, GREY, False, BODY, PP_ALIGN.CENTER)


def page2(prs):
    s = new_slide(prs)
    frame(s, 'Product Overview', 'SECTION 01')
    rect(s, 110, 185, 600, 250, fill=PANEL, line=BORDER)
    txt(s, 132, 203, 500, 28, 'What is AuditAI?', 15, PURPLE, True, HEAD)
    txt(s, 132, 242, 560, 180,
        'An AI-powered SaaS platform that audits smart\ncontracts in seconds — combining 500+ static\nrules, symbolic execution, and an LLM engine,\nwith A–F risk scoring and on-chain monitoring.',
        11, WHITE, False, BODY, spacing=1.4)
    rect(s, 730, 185, 610, 250, fill=PANEL, line=BORDER)
    txt(s, 752, 203, 500, 28, 'Why Now', 15, PURPLE, True, HEAD)
    stats = [('$3.8B+', 'lost to DeFi hacks'), ('$100B+', 'TVL at stake'), ('47 days', 'avg. to exploit')]
    for i, (big, sub) in enumerate(stats):
        txt(s, 752 + i * 205, 252, 200, 36, big, 18, YELLOW, True, HEAD)
        txt(s, 752 + i * 205, 296, 200, 20, sub, 9, GREY)
    txt(s, 752, 348, 570, 40, 'Audits cost $50K–$500K and take 4–8 weeks. We do it in 45 seconds.', 10, WHITE)
    chips = [('94.7%', 'DETECTION RATE'), ('45s', 'AVG. SCAN TIME'), ('6', 'CHAINS SUPPORTED'), ('500+', 'VULN PATTERNS')]
    for i, (big, sub) in enumerate(chips):
        x = 110 + i * 318
        rect(s, x, 470, 296, 140, fill=PANEL2, line=PURPLE)
        txt(s, x, 505, 296, 40, big, 21, WHITE, True, HEAD, PP_ALIGN.CENTER)
        txt(s, x, 560, 296, 22, sub, 9, GREY, False, BODY, PP_ALIGN.CENTER)
    rect(s, 110, 640, 1230, 84, fill=PANEL, line=BORDER)
    txt(s, 110, 640, 1230, 84, 'One-line pitch:  Ship secure contracts — before the hackers find the bugs.', 12, YELLOW, True, BODY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def page3(prs):
    s = new_slide(prs)
    frame(s, 'Problem Statement & Target Users', 'SECTION 02')
    rect(s, 110, 185, 1230, 200, fill=PANEL, line=RED)
    txt(s, 132, 203, 500, 28, 'The Problem', 15, RED, True, HEAD)
    txt(s, 132, 244, 1180, 130,
        '•  Manual audits take 4–8 weeks and cost $50K–$500K+ — critical bugs get exploited in as little as 6 hours.\n•  Even expert auditors miss 15–20% of vulnerabilities; static tools drown teams in false positives.\n•  Point-in-time audits go stale: 32% of exploits are reentrancy, 21% flash-loan attacks, 18% oracle manipulation.',
        11, WHITE, False, BODY, spacing=1.45)
    txt(s, 110, 408, 400, 20, 'TARGET USERS', 11, YELLOW, True, BODY)
    users = [('Protocol Teams', 'Instant CI/CD scans on every push & PR before deployment', 'Faster, safer shipping'),
             ('Security Auditors', 'Auto-triage findings; focus human effort on business logic', 'Higher coverage per hour'),
             ('Investors & DAOs', 'A–F risk score any contract before investing or voting', 'Data-driven decisions'),
             ('Exchanges & Launchpads', 'Screen listings at scale with the Risk Score API', 'Automated due diligence')]
    for i, (h, sub, gain) in enumerate(users):
        x = 110 + (i % 2) * 630
        y = 440 + (i // 2) * 145
        rect(s, x, y, 600, 125, fill=PANEL, line=BORDER)
        txt(s, x + 22, y + 16, 540, 26, h, 13, YELLOW, True, HEAD)
        txt(s, x + 22, y + 52, 556, 22, sub, 10, GREY)
        txt(s, x + 22, y + 84, 556, 20, '→  ' + gain, 9, GREEN)


def page4(prs):
    s = new_slide(prs)
    frame(s, 'Existing Solutions & Their Limitations', 'SECTION 03')
    sols = [('Slither', 'Static analyzer', 'High false positives; misses cross-contract & economic bugs'),
            ('Mythril', 'Symbolic executor', 'Slow on large contracts; shallow business-logic coverage'),
            ('Manual Audit Firms', 'Human experts', '$50K–$500K, 4–8 weeks, point-in-time only, hard to scale'),
            ('Basic LLM Checkers', 'AI chatbots', 'Generic, no exploit grounding, no monitoring, no CI/CD')]
    for i, (name, kind, lim) in enumerate(sols):
        x = 110 + (i % 2) * 630
        y = 190 + (i // 2) * 155
        rect(s, x, y, 600, 135, fill=PANEL, line=BORDER)
        txt(s, x + 22, y + 14, 400, 26, name, 14, WHITE, True, HEAD)
        txt(s, x + 22, y + 50, 400, 18, kind, 9, PURPLE, True, BODY)
        txt(s, x + 22, y + 80, 556, 40, lim, 10, GREY)
    rect(s, 110, 520, 1230, 200, fill=PANEL2, line=PURPLE)
    txt(s, 132, 538, 500, 28, 'The Gap AuditAI Fills', 15, PURPLE, True, HEAD)
    txt(s, 132, 578, 1180, 130,
        '•  Multi-engine depth of an audit firm — automated, repeatable, and 100x faster.\n•  AI semantic layer catches business-logic & economic attacks that pattern tools miss (+23% logic-vuln detection, −35% false positives).\n•  Continuous monitoring keeps verdicts fresh — no more point-in-time audit decay.\n•  API-first & CI/CD-native: security embedded in the dev workflow, not bolted on.',
        10, WHITE, False, BODY, spacing=1.4)


def page5(prs):
    s = new_slide(prs)
    frame(s, 'Proposed Solution & Key Innovations', 'SECTION 04')
    inno = [('01  Multi-Engine Analysis', 'Static rules (500+) + symbolic execution (10^6 paths) + AI semantic analysis — consensus scoring across engines'),
            ('02  AI Fix Suggestions', 'Context-aware patches with code, tests & deployment notes for every finding — not just an alert'),
            ('03  Continuous Monitoring', 'On-chain surveillance of deployed contracts with 2.3s alert latency across 6 chains'),
            ('04  Risk Score API', 'A–F grade for any contract address in one call — for investors, exchanges & DAOs')]
    for i, (h, sub) in enumerate(inno):
        x = 110 + (i % 2) * 630
        y = 195 + (i // 2) * 165
        rect(s, x, y, 600, 145, fill=PANEL, line=BORDER)
        txt(s, x + 22, y + 16, 540, 26, h, 13, YELLOW, True, HEAD)
        txt(s, x + 22, y + 54, 556, 70, sub, 10, WHITE, False, BODY, spacing=1.35)
    rect(s, 110, 545, 1230, 180, fill=PANEL, line=GREEN)
    txt(s, 132, 561, 500, 28, 'Competitive Edge', 15, GREEN, True, HEAD)
    edge = [('94.7%', 'vs 74.8% industry'), ('−35%', 'false positives'), ('45s', 'vs 4–8 weeks'), ('$29', 'vs $50K+ audits')]
    for i, (big, sub) in enumerate(edge):
        txt(s, 160 + i * 300, 605, 280, 38, big, 20, WHITE, True, HEAD)
        txt(s, 160 + i * 300, 650, 280, 20, sub, 9, GREY)


def page6(prs):
    s = new_slide(prs)
    frame(s, 'Technical Approach — System Workflow', 'SECTION 05')
    stages = [('INPUT', 'Paste code / .sol,.rs upload / on-chain address'),
              ('PARSE', 'Solidity • Vyper • Rust → unified AST + CFG'),
              ('STATIC', '500+ rules: reentrancy, access control, oracle bugs'),
              ('SYMBOLIC', 'Z3 constraint solving, 10^6 path exploration'),
              ('AI ENGINE', 'Fine-tuned LLM + RAG — logic & economic attacks'),
              ('AGGREGATE', 'Dedup, severity (CVSS), cross-engine consensus'),
              ('OUTPUT', 'Findings + AI fixes + A–F Risk Score + report'),
              ('MONITOR', 'Deployed-contract watch, 2.3s alerts, CI/CD hooks')]
    bw, bh = 276, 150
    for i, (h, sub) in enumerate(stages):
        x = 110 + (i % 4) * 300
        y = 200 + (i // 4) * 210
        rect(s, x, y, bw, bh, fill=PANEL, line=BORDER)
        txt(s, x + 16, y + 14, bw - 32, 24, h, 12, YELLOW, True, HEAD)
        txt(s, x + 16, y + 50, bw - 32, 90, '\n'.join(wrap(sub, 34)[:4]), 8.5, GREY, False, BODY, spacing=1.35)
        if i % 4 != 3:
            arrow(s, x + bw + 5, y + bh / 2 - 8, 14, 16)
    txt(s, 320, 578, 800, 24, 'Avg. full-pipeline scan: 45 seconds standard contract  •  3–8 min complex protocol', 11, GREEN, False, BODY, PP_ALIGN.CENTER)
    rect(s, 110, 620, 1230, 100, fill=PANEL2, line=PURPLE)
    txt(s, 132, 632, 500, 24, 'Detection Performance', 12, PURPLE, True, HEAD)
    perf = [('Reentrancy 99.2%', 'vs 87% industry'), ('Access Control 96.8%', 'vs 82%'),
            ('Flash Loan 94.1%', 'vs 76%'), ('Business Logic 89.3%', 'vs 58%')]
    for i, (a, b) in enumerate(perf):
        txt(s, 160 + i * 300, 668, 290, 20, a, 10, WHITE, True, BODY)
        txt(s, 160 + i * 300, 692, 290, 16, b, 8, GREY)


def page6b(prs):
    s = new_slide(prs)
    frame(s, 'Technology Stack — What Powers AuditAI', 'SECTION 06')
    stacks = [('FRONTEND', ['Next.js 15 + React 19 + TypeScript', 'Tailwind CSS 4 + Radix UI components', 'Three.js / GSAP visualizations', 'Lucide icon system']),
              ('BACKEND API', ['FastAPI (Python 3.12) · 14+ REST endpoints', 'JWT auth · API keys · Pydantic v2 schemas', 'Versioned /api/v1 · Uvicorn ASGI', 'reportlab PDF reports · Resend email']),
              ('DATA & CACHE', ['PostgreSQL 16 + SQLAlchemy ORM', 'Alembic migrations', 'Redis 7 cache + task broker', 'Celery background scan workers']),
              ('AI / ML ENGINE', ['Fine-tuned security LLM + RAG', 'Z3 symbolic constraint solver', '500+ static detection rules', 'CVSS severity scoring model']),
              ('BLOCKCHAIN LAYER', ['Solidity 0.4–0.8 · Vyper · Rust/Anchor', 'eth-account · Risk Score API (A–F)', '6 chains: ETH, Base, Arbitrum,', 'Polygon, BSC, Solana']),
              ('DEVOPS & CI/CD', ['GitHub Actions scan bot on PRs', 'Docker Compose infra (PG + Redis)', 'Vercel frontend deploys', 'Sentry monitoring + Resend email'])]
    for i, (h, lines) in enumerate(stacks):
        x = 110 + (i % 3) * 416
        y = 190 + (i // 3) * 218
        rect(s, x, y, 396, 198, fill=PANEL, line=BORDER)
        txt(s, x + 20, y + 14, 340, 24, h, 11, YELLOW, True, HEAD)
        txt(s, x + 20, y + 52, 360, 130, '\n'.join(lines), 8.5, WHITE, False, BODY, spacing=1.7)
    rect(s, 110, 630, 1230, 110, fill=PANEL2, line=PURPLE)
    txt(s, 132, 640, 500, 20, 'END-TO-END REQUEST FLOW', 9.5, PURPLE, True, BODY)
    flow = ['Next.js Dashboard', 'FastAPI REST', 'Celery + Redis Queue', 'Analysis Engines', 'PostgreSQL', 'Alerts + Reports']
    cw, sx = 172, 110 + (1230 - (6 * 172 + 5 * 26)) // 2
    for i, name in enumerate(flow):
        cx = sx + i * (cw + 26)
        rect(s, cx, 678, cw, 44, fill=PANEL, line=BORDER, lw=1.0)
        txt(s, cx, 678, cw, 44, name, 8.5, WHITE, False, BODY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < 5:
            arrow(s, cx + cw + 6, 692, 14, 16, fill=PURPLE)




def page_arch(prs):
    s = new_slide(prs)
    frame(s, 'System Architecture', 'SECTION 07')
    rows = [(190, 'CLIENT — Next.js 15 Dashboard (React 19, Tailwind, GSAP)',
             'Team seats, scan UI, remediation workflow, API console', BORDER),
            (292, 'JOB QUEUE — Redis 7 broker + Celery workers',
             'Async scan pipeline, retries, scheduled monitor sweeps', PURPLE),
            (394, 'DATA — PostgreSQL 16 (SQLAlchemy + Alembic)',
             'Orgs, scans, findings, monitors, usage meters, audit trails', GREEN)]
    rows_r = [(190, 'API GATEWAY — FastAPI (Uvicorn ASGI) + JWT / API keys',
               '/api/v1: scans, findings, monitoring, risk-score, billing', BORDER),
              (292, 'ANALYSIS ENGINES — Static + Symbolic (Z3) + AI LLM/RAG',
               '500+ rules · 10^6 paths · CVSS scoring · fix generation', PURPLE),
              (394, 'BLOCKCHAIN WATCHERS — eth-account RPC polling',
               'Ethereum, Base, Arbitrum, Polygon, BSC, Solana', GREEN)]
    for y, a, b, col in rows:
        rect(s, 110, y, 560, 74, fill=PANEL, line=col)
        txt(s, 130, y + 12, 540, 22, a, 10, YELLOW, True, HEAD)
        txt(s, 130, y + 40, 540, 18, b, 8.5, GREY)
    for y, a, b, col in rows_r:
        rect(s, 770, y, 560, 74, fill=PANEL, line=col)
        txt(s, 790, y + 12, 540, 22, a, 10, YELLOW, True, HEAD)
        txt(s, 790, y + 40, 540, 18, b, 8.5, GREY)
    for y in (264, 366):
        vline(s, 390, y, 28, BORDER, 2.0)
        vline(s, 1050, y, 28, BORDER, 2.0)
    for y in (222, 329, 431):
        hline(s, 670, y, 100, BORDER, 2.0)
    rect(s, 110, 496, 600, 60, fill=PANEL, line=YELLOW)
    txt(s, 130, 508, 570, 36, 'CI/CD: GitHub Action scans every push/PR with inline findings', 9.5, YELLOW, True, BODY, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 740, 496, 600, 60, fill=PANEL, line=YELLOW)
    txt(s, 760, 508, 570, 36, 'ALERTS: Slack / Discord / Email / Webhook (2.3s latency)', 9.5, YELLOW, True, BODY, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 110, 580, 1230, 160, fill=PANEL2, line=BORDER)
    txt(s, 132, 594, 500, 26, 'DEPLOYMENT & SCALE', 12, PURPLE, True, HEAD)
    dep = [('Vercel', 'frontend + edge'), ('Docker Compose', 'PG + Redis infra'),
           ('Worker autoscale', '10k+ scans/hour'), ('Sentry', 'error telemetry'),
           ('99.97%', 'API uptime, 12 mo')]
    for i, (a, b) in enumerate(dep):
        txt(s, 160 + i * 245, 636, 230, 28, a, 12, WHITE, True, HEAD)
        txt(s, 160 + i * 245, 668, 230, 18, b, 8.5, GREY)


def page_shots(prs):
    s = new_slide(prs)
    frame(s, 'Product Walkthrough — Live Dashboard', 'SECTION 08')
    txt(s, 110, 174, 900, 20, 'Real product UI · demo org data · scans, monitoring & alerting in production shape', 9.5, GREY)
    cards = [('scripts/out/shots/dashboard.png', 'Dashboard Overview',
              'Quick scan (code / address / GitHub), live usage meters, CI status'),
             ('scripts/out/shots/scans.png', 'Scans & Risk Grades',
              'A-F risk grades per contract — VulnerableVault graded F'),
             ('scripts/out/shots/monitoring.png', 'Continuous Monitoring',
              'On-chain watchtowers with health status + activity feed alerts')]
    for i, (img, title, sub) in enumerate(cards):
        x = 110 + i * 416
        rect(s, x, 208, 396, 380, fill=PANEL, line=BORDER)
        try:
            s.shapes.add_picture(img, IN(x + 10), IN(218), IN(376), IN(212))
        except Exception:
            txt(s, x + 20, 260, 300, 30, '[screenshot]', 12, GREY)
        txt(s, x + 18, 436, 360, 26, title, 12, YELLOW, True, HEAD)
        txt(s, x + 18, 466, 364, 80, '\n'.join(wrap(sub, 44)[:3]), 8.5, GREY, False, BODY, spacing=1.4)
    rect(s, 110, 620, 1230, 120, fill=PANEL2, line=PURPLE)
    txt(s, 132, 632, 400, 24, 'SHIPS TODAY', 11, PURPLE, True, HEAD)
    txt(s, 132, 660, 1180, 60, 'Next.js dashboard (11 pages) · FastAPI (14+ REST endpoints) · GitHub Action CI bot ·\nA-F Risk Score API · remediation workflow with sign-off & audit-trail export', 10, WHITE, False, BODY, spacing=1.3)

def page7(prs):
    s = new_slide(prs)
    frame(s, 'Tokenomics — SCRT Utility Token', 'SECTION 09')
    rect(s, 110, 185, 560, 260, fill=PANEL, line=BORDER)
    txt(s, 132, 201, 400, 28, 'Token Utility', 15, PURPLE, True, HEAD)
    utils = [('PAY', 'scans, monitoring, API & enterprise plans'),
             ('STAKE', 'discounts, priority access, higher API limits'),
             ('GOVERN', 'vote on features, fees, ecosystem fund'),
             ('EARN', 'bug bounties, referrals, community rewards')]
    for i, (a, b) in enumerate(utils):
        txt(s, 132, 246 + i * 46, 80, 22, a, 10, YELLOW, True, BODY)
        txt(s, 220, 246 + i * 46, 430, 22, b, 10, WHITE)
    rect(s, 700, 185, 640, 260, fill=PANEL, line=BORDER)
    txt(s, 722, 201, 400, 28, 'Token Distribution', 15, PURPLE, True, HEAD)
    dist = [('Ecosystem Fund', 35, BORDER), ('Community Rewards', 25, PURPLE),
            ('Team & Advisors', 20, YELLOW), ('Private Sale', 10, GREEN),
            ('Public Sale', 5, ORANGE), ('Liquidity', 5, RED)]
    for i, (name, pct, col) in enumerate(dist):
        y = 248 + i * 31
        txt(s, 722, y, 170, 18, name, 9, WHITE)
        rect(s, 900, y + 2, 6.0 * pct, 12, fill=col, line=None, rad=0.5)
        txt(s, 1240, y, 85, 18, f'{pct}%', 9, WHITE, True, BODY, PP_ALIGN.RIGHT)
    txt(s, 722, 428, 600, 16, 'Vesting: Team 1y cliff / 3y linear • Private 6m cliff / 1y • Public 25% TGE', 7.5, GREY)
    rect(s, 110, 468, 560, 272, fill=PANEL, line=GREEN)
    txt(s, 132, 484, 480, 28, 'Economic Model — The Flywheel', 15, GREEN, True, HEAD)
    fly = ['1. Demand: teams need constant security coverage',
           '2. Utility: SCRT pays for scans, staking, governance',
           '3. Value accrual: fee burning + revenue buybacks',
           '4. Growth: rewards attract auditors & contributors',
           '5. Platform improves → more demand → repeat']
    txt(s, 132, 528, 520, 200, '\n'.join(fly), 10, WHITE, False, BODY, spacing=1.5)
    rect(s, 700, 468, 640, 272, fill=PANEL, line=RED)
    txt(s, 722, 484, 480, 28, 'Deflationary Mechanisms', 15, RED, True, HEAD)
    defs = [('Fee Burning', 'share of every SCRT service fee burned forever'),
            ('Buyback & Burn', 'fiat revenue used to buy & burn SCRT'),
            ('Staking Lockups', 'loyalty tiers reduce circulating supply'),
            ('Bounty Escrow', 'treasury SCRT reserved for bounties & grants')]
    for i, (a, b) in enumerate(defs):
        txt(s, 722, 532 + i * 52, 300, 22, '• ' + a, 10, YELLOW, True, BODY)
        txt(s, 742, 556 + i * 52, 560, 18, b, 8.5, GREY)


def page8(prs):
    s = new_slide(prs)
    frame(s, 'Business Model & Pricing', 'SECTION 10')
    plans = [('FREE', '$0/mo', ['50 scans/mo', '1 monitored contract', 'GitHub Action', 'Community support'], GREY),
             ('PRO', '$29/mo', ['500 scans/mo', '10 monitored contracts', 'AI fix suggestions', 'Slack/Discord alerts'], BORDER),
             ('TEAM', '$99/mo', ['2,000 scans/mo', '50 monitored contracts', 'Remediation workflow', 'Unlimited seats'], PURPLE),
             ('ENTERPRISE', 'Custom', ['Unlimited scans', 'Risk Score API', '99.95% SLA', 'Dedicated support'], GREEN)]
    for i, (name, price, feats, col) in enumerate(plans):
        x = 110 + i * 318
        rect(s, x, 190, 296, 330, fill=PANEL, line=col, lw=1.6)
        txt(s, x, 216, 296, 28, name, 14, col if col != GREY else WHITE, True, HEAD, PP_ALIGN.CENTER)
        txt(s, x, 250, 296, 36, price, 18, WHITE, True, HEAD, PP_ALIGN.CENTER)
        for j, ft in enumerate(feats):
            txt(s, x + 30, 318 + j * 34, 260, 20, '• ' + ft, 9, GREY)
        if i == 1:
            txt(s, x, 170, 296, 16, 'MOST POPULAR', 8, YELLOW, True, BODY, PP_ALIGN.CENTER)
    rect(s, 110, 545, 1230, 195, fill=PANEL, line=BORDER)
    txt(s, 132, 560, 500, 28, 'Revenue Streams & Unit Economics', 15, PURPLE, True, HEAD)
    rev = [('Subscription SaaS', 'core recurring revenue across 4 tiers'),
           ('Usage-based', '$0.10 per extra scan; volume discounts'),
           ('API Licensing', 'Risk Score API for exchanges & wallets'),
           ('Enterprise', 'custom infra, compliance docs, SLAs')]
    for i, (a, b) in enumerate(rev):
        txt(s, 160 + (i % 2) * 610, 598 + (i // 2) * 64, 400, 22, '• ' + a, 10, YELLOW, True, BODY)
        txt(s, 180 + (i % 2) * 610, 620 + (i // 2) * 64, 480, 18, b, 9, GREY)


def page9(prs):
    s = new_slide(prs)
    frame(s, 'Development Roadmap', 'SECTION 11')
    hline(s, 140, 258, 1170, BORDER, 3.0)
    phases = [('Q3 2026', ['Cross-chain invariant monitoring', 'Advanced fuzzing engine', 'Formal verification module']),
              ('Q4 2026', ['MEV protection analysis', 'Gas optimization recommender', 'Custom rule creation SDK']),
              ('Q1 2027', ['ZK-proof verification', 'Move language support', 'Automated exploit simulation']),
              ('Q2 2027', ['WASM contract support', 'AI exploit prediction', 'Decentralized audit marketplace'])]
    for i, (q, items) in enumerate(phases):
        x = 110 + i * 318
        dot(s, x + 139, 250, 18, 18)
        rect(s, x, 300, 296, 280, fill=PANEL, line=BORDER)
        txt(s, x + 20, 316, 250, 26, q, 14, YELLOW, True, HEAD)
        txt(s, x + 20, 360, 256, 200, '\n'.join('• ' + it for it in items), 9.5, WHITE, False, BODY, spacing=1.8)
    rect(s, 110, 610, 1230, 130, fill=PANEL, line=GREEN)
    txt(s, 132, 624, 500, 28, 'Long-Term Vision', 15, GREEN, True, HEAD)
    txt(s, 132, 662, 1180, 60, 'A fully automated security layer for the entire blockchain ecosystem — every contract continuously\nmonitored and protected against emerging threats in real time.', 10.5, WHITE, False, BODY, spacing=1.35)


def page10(prs):
    s = new_slide(prs)
    frame(s, 'Team & Why We Win', 'SECTION 12')
    txt(s, 110, 182, 500, 20, 'TEAM UNFILTERED DEVS', 11, YELLOW, True, BODY)
    team = [('Ajay Yadav', 'Team Lead & Frontend'), ('Rangesh Gupta', 'AI/ML & Backend'),
            ('Shivam Prajapati', 'UI/UX & Frontend'), ('Poornita Sahu', 'Research & Presentation')]
    for i, (n, r) in enumerate(team):
        x = 110 + i * 318
        rect(s, x, 220, 296, 130, fill=PANEL, line=BORDER)
        dot(s, x + 118, 238, 60, 60, fill=PANEL2, line=PURPLE, lw=1.6)
        txt(s, x + 118, 238, 60, 60, n[0], 16, PURPLE, True, HEAD, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(s, x, 306, 296, 22, n, 11, WHITE, True, BODY, PP_ALIGN.CENTER)
        txt(s, x, 328, 296, 16, r, 8, GREY, False, BODY, PP_ALIGN.CENTER)
    rect(s, 110, 385, 1230, 165, fill=PANEL, line=BORDER)
    txt(s, 132, 400, 500, 28, 'Why We Win', 15, PURPLE, True, HEAD)
    txt(s, 132, 442, 1180, 100,
        '•  Only platform combining static + symbolic + AI engines with continuous monitoring at SaaS pricing\n•  Working product: Next.js dashboard, FastAPI backend, GitHub Action CI/CD — already scanning contracts\n•  AI fix suggestions close the loop from detection to remediation — a category first at this price point',
        10, WHITE, False, BODY, spacing=1.45)
    rect(s, 110, 580, 1230, 145, fill=PANEL2, line=YELLOW)
    txt(s, 132, 596, 500, 28, 'The Ask', 15, YELLOW, True, HEAD)
    txt(s, 132, 638, 1180, 24, 'Raising to scale go-to-market, expand chain coverage, and grow the AI detection research team.', 11, WHITE)
    txt(s, 132, 672, 1180, 20, 'Contact: team@auditai.io  •  auditai.io  •  github.com/securithm', 9, GREY)


def page11(prs):
    s = new_slide(prs)
    frame(s, 'References & Disclosure of Existing Work', 'SECTION 13')
    refs = [('01', 'SWC Registry', 'Smart Contract Weakness Classification', 'swcregistry.io'),
            ('02', 'OWASP Smart Contract Top 10', 'Industry vulnerability taxonomy', 'owasp.org/www-project-smart-contract-top-10'),
            ('03', 'Slither', 'Static analysis framework reference', 'github.com/crytic/slither'),
            ('04', 'Mythril', 'Symbolic execution approach reference', 'github.com/Consensys/mythril'),
            ('05', 'Rekt News', 'DeFi attack database & case studies', 'rekt.news'),
            ('06', 'Certora', 'Formal verification of smart contracts', 'certora.com'),
            ('07', 'ZEUS / S-gram Research', 'Academic smart-contract analysis papers', 'IEEE S&P / ACM'),
            ('08', 'Securithm Whitepaper v1.0', 'Full technical methodology document', 'securithm.com/whitepaper')]
    for i, (num, name, desc, url) in enumerate(refs):
        x = 110 + (i % 2) * 630
        y = 195 + (i // 2) * 130
        rect(s, x, y, 600, 110, fill=PANEL, line=BORDER)
        txt(s, x + 22, y + 18, 50, 22, num, 10, PURPLE, True, BODY)
        txt(s, x + 70, y + 16, 500, 24, name, 12, WHITE, True, HEAD)
        txt(s, x + 70, y + 46, 500, 18, desc, 8.5, GREY)
        txt(s, x + 70, y + 72, 500, 18, url, 8.5, GREEN)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for pg in [page1, page2, page3, page4, page5, page6, page6b, page_arch, page_shots, page7, page8, page9, page10, page11]:
        pg(prs)
    core = prs.core_properties
    core.title = 'AuditAI — AI-Powered Smart Contract Security Platform'
    core.author = 'Team Unfiltered Devs'
    core.subject = 'Securithm SaaS Audit AI — Investor Pitch Deck'
    prs.save('AuditAI_Investor_Pitch_Deck.pptx')
    print('PPTX written: AuditAI_Investor_Pitch_Deck.pptx')


if __name__ == '__main__':
    main()
